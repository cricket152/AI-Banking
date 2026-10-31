"""按 ppt-kit 逐页指南生成 AI Banking 答辩 PPT：10 页正文 + 2 页备选附录。

配文、讲稿与版式依据 ../ppt-kit/slides.json 与 ../ppt-kit/ppt-production-guide.docx，
图像素材全部来自 ../ppt-kit/assets（14 张运行原图、4 张细节裁切、3 组 PNG/SVG 架构图）。
文本与流程节点均为可编辑元素；截图只做等比裁切与取景，不改动界面文字与结果。
"""

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
KIT = ROOT / "ppt-kit" / "assets"
OUT = Path(__file__).resolve().parent / "AI-Banking-defense.pptx"

FONT_EA = "Microsoft YaHei"
FONT_LATIN = "Arial"

COLORS = {
    "paper": "F7F4ED",
    "ink": "20323D",
    "muted": "5E6E73",
    "line": "DED6C7",
    "white": "FFFDF9",
    "gold": "9C6A2F",
    "gold_pale": "F5E9D7",
    "teal": "2F7466",
    "teal_pale": "E4EFE9",
    "blue_pale": "E8EEF1",
    "red": "A8524C",
    "red_pale": "F5E9E7",
}

SLIDE_W = 13.333
SLIDE_H = 7.5
M = 0.56
RIGHT = SLIDE_W - M
CONTENT_W = RIGHT - M
TOTAL_PAGES = 10


def rgb(name):
    return RGBColor.from_string(COLORS.get(name, name))


def style_run(run, size, color="ink", bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    run.font.name = FONT_LATIN
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        rpr.append(rpr.makeelement(qn(tag), {"typeface": FONT_EA}))


def para(runs, align=PP_ALIGN.LEFT, line_spacing=1.25, space_before=0, space_after=0):
    """runs: [(text, size, color, bold), ...]"""
    if isinstance(runs, tuple):
        runs = [runs]
    return {
        "runs": runs,
        "align": align,
        "line_spacing": line_spacing,
        "space_before": space_before,
        "space_after": space_after,
    }


def add_paras(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = anchor
    for index, spec in enumerate(paras):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = spec["align"]
        paragraph.line_spacing = spec["line_spacing"]
        paragraph.space_before = Pt(spec["space_before"])
        paragraph.space_after = Pt(spec["space_after"])
        for text, size, color, bold in spec["runs"]:
            run = paragraph.add_run()
            run.text = text
            style_run(run, size, color, bold)
    return box


def add_rect(slide, x, y, w, h, fill="white", line=None, line_w=0.75, radius=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        shape.adjustments[0] = radius
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_w)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def add_oval(slide, x, y, w, h, line="teal", width=1.0):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.background()
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(width)
    shape.shadow.inherit = False
    return shape


def add_line(slide, x1, y1, x2, y2, color="line", width=0.75):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    connector.line.color.rgb = rgb(color)
    connector.line.width = Pt(width)
    connector.shadow.inherit = False
    return connector


def add_arrow(slide, cx, cy, direction="right", size=0.17, color="gold"):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ISOSCELES_TRIANGLE,
        Inches(cx - size / 2),
        Inches(cy - size / 2),
        Inches(size),
        Inches(size),
    )
    shape.rotation = {"right": 90, "down": 180, "left": 270, "up": 0}[direction]
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def set_alpha(picture, percent):
    blip = picture._element.blipFill.find(qn("a:blip"))
    blip.append(blip.makeelement(qn("a:alphaModFix"), {"amt": str(int(percent * 1000))}))


def add_picture(slide, path, x, y, w, h, crop=None, alpha=None):
    """等比放入 (x, y, w, h) 框内；crop = (left, right, top, bottom) 比例裁切。"""
    path = Path(path)
    with Image.open(path) as source:
        iw, ih = source.size
    left, right, top, bottom = crop or (0, 0, 0, 0)
    vw = iw * (1 - left - right)
    vh = ih * (1 - top - bottom)
    scale = min(w / vw, h / vh)
    pw, ph = vw * scale, vh * scale
    px, py = x + (w - pw) / 2, y + (h - ph) / 2
    picture = slide.shapes.add_picture(
        str(path), Inches(px), Inches(py), Inches(pw), Inches(ph)
    )
    if crop:
        picture.crop_left = left
        picture.crop_right = right
        picture.crop_top = top
        picture.crop_bottom = bottom
    if alpha is not None:
        set_alpha(picture, alpha)
    picture.shadow.inherit = False
    return picture, (px, py, pw, ph)


def add_framed_picture(slide, path, x, y, w, h, crop=None, alpha=None, line="line", line_w=0.75):
    """等比放入框内并按图片实际范围描边，避免框与图不等宽。"""
    picture, (px, py, pw, ph) = add_picture(slide, path, x, y, w, h, crop=crop, alpha=alpha)
    add_rect(slide, px, py, pw, ph, None, line, line_w)
    return picture


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb("paper")
    return slide


def header(slide, kicker, title, subtitle=None):
    add_paras(
        slide,
        M,
        0.44,
        8.0,
        0.26,
        [para((kicker, 12, "teal", True), line_spacing=1.0)],
    )
    add_paras(
        slide,
        M,
        0.74,
        CONTENT_W,
        0.66,
        [para((title, 32, "ink", True), line_spacing=1.05)],
    )
    if subtitle:
        add_paras(
            slide,
            M,
            1.44,
            CONTENT_W,
            0.30,
            [para((subtitle, 14, "muted", False), line_spacing=1.2)],
        )
    add_line(slide, M, 1.80, RIGHT, 1.80, "line", 0.75)


def footer(slide, label, note="AI Banking · 本地模拟 · 虚构数据 · 脚本化 Agent 演示"):
    add_line(slide, M, 7.00, RIGHT, 7.00, "line", 0.75)
    add_paras(slide, M, 7.08, 8.6, 0.22, [para((note, 9, "muted", False), line_spacing=1.0)])
    add_paras(
        slide,
        RIGHT - 2.2,
        7.08,
        2.2,
        0.22,
        [para((label, 9, "muted", False), align=PP_ALIGN.RIGHT, line_spacing=1.0)],
    )


def tag(slide, x, y, text, fill="teal_pale", color="teal", size=11, pad=0.16):
    width = pad * 2 + len(text) * (size / 72) * 0.92
    add_rect(slide, x, y, width, 0.30, fill, None, radius=0.4)
    add_paras(
        slide,
        x,
        y + 0.055,
        width,
        0.20,
        [para((text, size, color, True), align=PP_ALIGN.CENTER, line_spacing=1.0)],
    )
    return width


def caption(slide, x, y, w, text, align=PP_ALIGN.LEFT, size=12.5):
    add_paras(
        slide,
        x,
        y,
        w,
        0.26,
        [para((text, size, "muted", False), align=align, line_spacing=1.15)],
    )


def bullet_row(slide, x, y, w, keyword, body, size_key=21, size_body=16.5):
    add_paras(
        slide,
        x,
        y,
        2.0,
        0.36,
        [para((keyword, size_key, "ink", True), line_spacing=1.1)],
    )
    add_paras(
        slide,
        x,
        y + 0.42,
        w,
        0.52,
        [para((body, size_body, "ink", False), line_spacing=1.25)],
    )


def page_01(prs):
    slide = new_slide(prs)
    add_rect(slide, M, 0.50, 0.46, 0.46, "gold")
    add_paras(
        slide,
        M,
        0.60,
        0.46,
        0.30,
        [para(("A", 20, "white", True), align=PP_ALIGN.CENTER, line_spacing=1.0)],
    )
    add_paras(
        slide,
        1.16,
        0.50,
        4.6,
        0.28,
        [para(("AI-BANKING", 15, "ink", True), line_spacing=1.0)],
    )
    add_paras(
        slide,
        1.16,
        0.76,
        5.4,
        0.24,
        [para(("本地模拟银行原型 · 独立虚构数据", 11, "muted", False), line_spacing=1.0)],
    )
    add_paras(
        slide,
        M,
        1.16,
        CONTENT_W,
        0.80,
        [para(("AI Banking 让金融意图有依据 让执行有边界", 34, "ink", True), line_spacing=1.05)],
    )
    add_line(slide, M, 2.04, M + 1.0, 2.04, "gold", 2.5)

    add_paras(
        slide,
        M,
        2.24,
        4.35,
        0.60,
        [para(("一个可以实际运行的个人模拟银行助手", 18, "gold", True), line_spacing=1.2)],
    )
    bullets = [
        "账单分析、订阅管理与受控转账的一体化工作台",
        "自然语言提出需求，后端核验权限与资金事实",
        "本地模拟银行原型",
    ]
    y = 3.06
    for item in bullets:
        add_rect(slide, M + 0.02, y + 0.11, 0.16, 0.03, "gold")
        add_paras(
            slide,
            M + 0.32,
            y,
            4.05,
            0.60,
            [para((item, 16.5, "ink", False), line_spacing=1.3)],
        )
        y += 0.78
    add_paras(
        slide,
        M,
        5.62,
        4.35,
        0.60,
        [
            para(("10 页正文 · 预计讲解 8 分 40 秒", 13, "muted", False), line_spacing=1.3),
            para(
                (
                    "文档版本 v1.2 · 功能基线 9389868 · 整理日期 2026-10-31",
                    11.5,
                    "muted",
                    False,
                ),
                line_spacing=1.3,
                space_before=4,
            ),
        ],
    )

    add_framed_picture(slide, KIT / "screenshots" / "02-overview.png", 5.30, 2.10, 7.47, 4.30)
    caption(slide, 5.30, 6.52, 7.47, "个人财务概览 · 独立虚构数据", PP_ALIGN.CENTER)
    footer(slide, "01 / 10")


def page_02(prs):
    slide = new_slide(prs)
    header(
        slide,
        "02 / 用户问题",
        "金融代办需要把对象和结果核对清楚",
        "查询结果、操作确认与执行回执分别呈现，模型文字不替代银行状态。",
    )
    rows = [
        ("看账单", "要能展开原始流水与判断依据"),
        ("找对象", "同名收款人需要备注和尾号澄清"),
        ("做操作", "批准内容与最终回执分开核对"),
    ]
    y = 2.24
    for index, (keyword, body) in enumerate(rows):
        if index == 0:
            add_rect(slide, 0.62, y + 0.14, 0.46, 0.46, None, "teal", 1.0)
            add_line(slide, 0.74, y + 0.26, 0.96, y + 0.26, "teal", 1.0)
            add_line(slide, 0.74, y + 0.36, 0.96, y + 0.36, "teal", 1.0)
            add_line(slide, 0.74, y + 0.46, 0.88, y + 0.46, "teal", 1.0)
        elif index == 1:
            add_oval(slide, 0.60, y + 0.18, 0.32, 0.32, "gold", 1.0)
            add_oval(slide, 0.80, y + 0.18, 0.32, 0.32, "gold", 1.0)
        else:
            add_rect(slide, 0.62, y + 0.14, 0.46, 0.46, None, "ink", 1.0, radius=0.12)
            add_line(slide, 0.74, y + 0.37, 0.84, y + 0.47, "ink", 1.0)
            add_line(slide, 0.84, y + 0.47, 0.96, y + 0.27, "ink", 1.0)
        bullet_row(slide, 1.32, y + 0.06, 4.9, keyword, body)
        if index < len(rows) - 1:
            add_line(slide, 1.32, y + 1.02, 6.20, y + 1.02, "line", 0.75)
        y += 1.26

    crop = (300 / 1440, (1440 - 950) / 1440, 30 / 1000, (1000 - 480) / 1000)
    add_framed_picture(slide, KIT / "screenshots" / "06-payees.png", 6.72, 2.20, 6.05, 4.05, crop=crop)
    caption(slide, 6.72, 6.36, 6.05, "同名收款人以备注和尾号区分 · 界面截图", PP_ALIGN.CENTER)
    footer(slide, "02 / 10")


def page_03(prs):
    slide = new_slide(prs)
    header(
        slide,
        "03 / 能力地图",
        "一个工作台连接六类模拟银行场景",
        "区分产品能力与模型权限：Agent 查询或准备有限意图，写操作由界面发起。",
    )
    cards = [
        ("账单洞察", ["查询"], "本期与前一等长区间同口径汇总，异常线索可回到原始流水"),
        ("订阅管理", ["查询", "意图准备"], "识别周期扣费候选，用户核对后登记；取消需逐步确认"),
        ("转账", ["查询", "意图准备"], "准备对象与金额的转账意图，执行前复核额度与余额"),
        ("卡片", ["查询", "界面办理"], "Agent 只读脱敏卡信息；申请与调额需密码再认证"),
        ("理财", ["查询", "界面办理"], "模拟产品目录只读；申购与赎回由界面逐笔确认"),
        ("日期计划", ["查询", "界面办理"], "保存日期与预算，用户手动下单；Agent 不可创建或下单"),
    ]
    fills = {"查询": "teal_pale", "意图准备": "gold_pale", "界面办理": "blue_pale"}
    inks = {"查询": "teal", "意图准备": "gold", "界面办理": "ink"}
    positions = [(M, 2.12), (4.72, 2.12), (8.88, 2.12), (M, 3.92), (4.72, 3.92), (8.88, 3.92)]
    for (name, entries, body), (x, y) in zip(cards, positions):
        add_rect(slide, x, y, 3.89, 1.66, "white", "line", 0.75, radius=0.05)
        add_paras(
            slide,
            x + 0.22,
            y + 0.18,
            3.45,
            0.32,
            [para((name, 19, "ink", True), line_spacing=1.0)],
        )
        add_paras(
            slide,
            x + 0.22,
            y + 0.60,
            3.45,
            0.62,
            [para((body, 13.5, "muted", False), line_spacing=1.3)],
        )
        offset = x + 0.22
        for entry in entries:
            offset += tag(slide, offset, y + 1.20, entry, fills[entry], inks[entry]) + 0.12

    add_rect(slide, M, 5.78, CONTENT_W, 0.86, "white", "line", 0.75, radius=0.05)
    add_paras(
        slide,
        M + 0.24,
        5.92,
        CONTENT_W - 0.48,
        0.60,
        [
            para(
                (
                    "Agent 通过 12 个类型化工具查询数据或准备有限意图，写操作入口与后端能力一致",
                    14,
                    "ink",
                    True,
                ),
                line_spacing=1.1,
            ),
            para(
                (
                    "list_accounts · get_transactions · get_billing_summary · get_billing_entries · "
                    "list_payees · find_payees · list_subscriptions · prepare_subscription_cancellation · "
                    "prepare_transfer · get_operation · list_wealth_products · list_cards",
                    10.5,
                    "muted",
                    False,
                ),
                line_spacing=1.2,
                space_before=4,
            ),
        ],
    )
    footer(slide, "03 / 10")


def page_04(prs):
    slide = new_slide(prs)
    header(
        slide,
        "04 / 系统架构",
        "模型提出意图 后端逐次复核再执行",
        "原生 Web + FastAPI + SQLite，单进程内的职责模块；安全检查由确定性代码实施，没有额外审核模型。",
    )
    steps = [
        ("浏览器 · 同源 Web", "用户问题、界面选择、任务与操作进度"),
        ("FastAPI 身份入口（A）", "会话身份 · Origin/Host · CSRF（写请求）· 输入大小与格式"),
        ("每轮上下文装配", "固定系统约束 + 状态 + 最新问题；历史超 12 轮 → 有界摘录"),
        ("单主 Agent / 模型工具循环", "理解问题 → 选择工具 → 读取返回，可继续调用、澄清或回答"),
        ("B 每次工具调用前的安全校验", "会话有效性 · 工具白名单 · Pydantic 参数 · 本人对象归属"),
        ("C 意图准备 / 任务约束", "转账、取消订阅；固定方案的对象 / 金额 / 顺序 / 幂等"),
        ("D 批准与执行前再次复核", "逐笔挑战或用户指定预授权 → 额度 / 余额 / 期限 → SQLite 事务"),
        ("回执", "双边记账、原子转账；结果回到 Agent，最终以操作与账本为准"),
    ]
    top = 2.04
    row = 0.565
    add_line(slide, 0.76, top + 0.30, 0.76, top + row * (len(steps) - 1) + 0.30, "line", 1.0)
    for index, (title, detail) in enumerate(steps):
        y = top + index * row
        badge = "gold" if index in (1, 4, 5, 6) else "teal"
        add_rect(slide, 0.61, y + 0.05, 0.30, 0.30, "paper", badge, 1.1, radius=0.5)
        add_paras(
            slide,
            0.61,
            y + 0.11,
            0.30,
            0.20,
            [para((str(index + 1), 10.5, badge, True), align=PP_ALIGN.CENTER, line_spacing=1.0)],
        )
        add_paras(
            slide,
            1.06,
            y + 0.08,
            7.50,
            0.44,
            [
                para(
                    [
                        (title, 14.5, "ink", True),
                        ("  ·  " + detail, 12, "muted", False),
                    ],
                    line_spacing=1.2,
                )
            ],
        )

    checks = [
        ("A", "请求入口", "会话身份、Origin/Host、CSRF、输入长度与格式"),
        ("B", "每次工具调用", "工具白名单、参数 Schema、本人数据归属"),
        ("C", "意图与任务约束", "固定步骤、对象与金额、状态与幂等"),
        ("D", "批准与执行", "逐笔确认或有效预授权；SQLite 事务复核"),
    ]
    y = 2.04
    for letter, title, detail in checks:
        add_rect(slide, 9.00, y, 3.77, 1.02, "white", "line", 0.75, radius=0.06)
        add_rect(slide, 9.18, y + 0.20, 0.42, 0.42, "gold_pale", None, radius=0.5)
        add_paras(
            slide,
            9.18,
            y + 0.28,
            0.42,
            0.24,
            [para((letter, 13, "gold", True), align=PP_ALIGN.CENTER, line_spacing=1.0)],
        )
        add_paras(
            slide,
            9.74,
            y + 0.18,
            2.85,
            0.26,
            [para((title, 15, "ink", True), line_spacing=1.0)],
        )
        add_paras(
            slide,
            9.74,
            y + 0.50,
            2.85,
            0.42,
            [para((detail, 11.5, "muted", False), line_spacing=1.2)],
        )
        y += 1.14
    caption(
        slide,
        9.00,
        6.62,
        3.77,
        "运行图按 agent-runtime.svg 重绘为可编辑元素",
        size=11,
    )
    footer(slide, "04 / 10")


def page_05(prs):
    slide = new_slide(prs)
    header(
        slide,
        "05 / 账单洞察",
        "每条异常线索都能回到原始流水",
        "规则提示是可解释的核对线索，不是欺诈判定。",
    )
    rules = [
        (
            2.12,
            1.36,
            "规则一 · 固定大额",
            "单笔支出不少于 ¥1,000 即提示人工核对",
            "示例：演示大额购物 ¥1,500.00",
        ),
        (
            3.62,
            1.92,
            "规则二 · 相对历史中位数",
            "本期单笔超过前一等长区间同分类中位数的 3 倍，且不少于 ¥100",
            "示例：餐饮 ¥200.00 · 前期中位数 ¥20.00",
        ),
    ]
    for y, h, title, body, sample in rules:
        add_rect(slide, M, y, 4.42, h, "white", "line", 0.75, radius=0.05)
        add_rect(slide, M, y, 0.05, h, "gold")
        add_paras(
            slide,
            M + 0.26,
            y + 0.18,
            4.0,
            0.30,
            [para((title, 17, "ink", True), line_spacing=1.0)],
        )
        add_paras(
            slide,
            M + 0.26,
            y + 0.56,
            4.0,
            0.56,
            [para((body, 14.5, "ink", False), line_spacing=1.3)],
        )
        add_paras(
            slide,
            M + 0.26,
            y + h - 0.42,
            4.0,
            0.28,
            [para((sample, 13, "gold", True), line_spacing=1.1)],
        )
    add_arrow(slide, 5.12, 4.30, "right", 0.20, "gold")
    add_paras(
        slide,
        M,
        5.72,
        4.42,
        0.80,
        [
            para(
                (
                    "同一事务读取本期与前一等长区间；消费 / 退款 / 转账分列，支持按原区间分页核对",
                    13.5,
                    "muted",
                    False,
                ),
                line_spacing=1.3,
            )
        ],
    )
    add_framed_picture(slide, KIT / "details" / "bill-evidence.png", 5.42, 2.12, 7.35, 3.74)
    caption(
        slide,
        5.42,
        6.02,
        7.35,
        "合成流水示例 · 异常线索用于人工核对；截图原文保留 ¥1,500 与 ¥200 两条命中",
        PP_ALIGN.CENTER,
    )
    footer(slide, "05 / 10")


def page_06(prs):
    slide = new_slide(prs)
    header(
        slide,
        "06 / 核心业务演示",
        "从订阅证据到两步执行形成闭环",
        "识别结果只是候选；登记与执行都需要用户显式确认。",
    )
    steps = [
        ("① 证据识别", "同账户同商户三次以上等额月度支出 → 候选"),
        ("② 用户核对并登记", "登记不扣款，也不产生授权"),
        ("③ 预览并保存方案", "对象、金额与步骤顺序；保存本身不付款"),
        ("④ 逐步批准执行", "每步独立确认，成功后才继续下一步"),
    ]
    x = M
    for index, (title, body) in enumerate(steps):
        add_rect(slide, x, 2.10, 2.82, 1.02, "white", "line", 0.75, radius=0.05)
        add_paras(
            slide,
            x + 0.20,
            2.26,
            2.45,
            0.28,
            [para((title, 15.5, "ink", True), line_spacing=1.0)],
        )
        add_paras(
            slide,
            x + 0.20,
            2.60,
            2.45,
            0.44,
            [para((body, 12, "muted", False), line_spacing=1.25)],
        )
        if index < 3:
            add_arrow(slide, x + 2.97, 2.61, "right", 0.16, "gold")
        x += 3.11

    add_framed_picture(slide, KIT / "details" / "subscription-evidence.png", M, 3.36, 2.62, 3.06)
    caption(slide, M, 6.46, 2.62, "订阅证据 · 周期扣费核对", PP_ALIGN.CENTER, 12)

    add_framed_picture(slide, KIT / "details" / "workflow-complete.png", 3.36, 3.36, 4.90, 3.06)

    zoom_crop = (0.0, (1080 - 300) / 1080, 228 / 710, (710 - 256) / 710)
    add_framed_picture(
        slide,
        KIT / "details" / "workflow-complete.png",
        3.36,
        6.42,
        3.20,
        0.42,
        crop=zoom_crop,
        line="gold",
    )
    caption(slide, 6.66, 6.46, 5.60, "放大标注：任务进度 · 已完成 2/2 步（截图原文）", size=12)

    add_rect(slide, 8.48, 3.36, 4.29, 3.06, "gold_pale", None, radius=0.04)
    add_paras(
        slide,
        8.72,
        3.62,
        3.85,
        1.30,
        [
            para(("2 / 2 步已完成", 22, "ink", True), line_spacing=1.1),
            para(
                (
                    "回执来自实际模拟执行器；界面以任务步骤记录真实状态。",
                    13.5,
                    "ink",
                    False,
                ),
                line_spacing=1.3,
                space_before=8,
            ),
        ],
    )
    add_line(slide, 8.72, 5.14, 12.53, 5.14, "gold", 0.75)
    add_paras(
        slide,
        8.72,
        5.30,
        3.85,
        0.90,
        [
            para(("登记不扣款 · 取消不退款", 13.5, "ink", True), line_spacing=1.2),
            para(
                (
                    "未来减少的订阅支出不会被当作当前到账资金。",
                    13,
                    "muted",
                    False,
                ),
                line_spacing=1.3,
                space_before=6,
            ),
        ],
    )
    footer(slide, "06 / 10")


def page_07(prs):
    slide = new_slide(prs)
    header(
        slide,
        "07 / 扩展场景",
        "日期预算连接计划与用户确认",
        "计划只保存日期与预算；创建时校验余额，但不冻结或预留资金。",
    )
    nodes = [
        ("① 日期与预算", "生日或纪念日场景"),
        ("② 余额校验", "按当时可用余额校验"),
        ("③ 保存计划", "不冻结、不预留资金"),
        ("④ 用户手动下单", "可操作时间内点击界面按钮"),
        ("⑤ 待确认操作", "生成待确认的模拟转账"),
        ("⑥ 后端回执", "与原操作状态联动"),
    ]
    top_row = [(M, 2.12), (4.72, 2.12), (8.88, 2.12)]
    bottom_row = [(8.88, 3.38), (4.72, 3.38), (M, 3.38)]
    for (title, body), (x, y) in zip(nodes, top_row + bottom_row):
        add_rect(slide, x, y, 3.89, 1.02, "white", "line", 0.75, radius=0.05)
        add_paras(
            slide,
            x + 0.22,
            y + 0.18,
            3.45,
            0.28,
            [para((title, 17, "ink", True), line_spacing=1.0)],
        )
        add_paras(
            slide,
            x + 0.22,
            y + 0.54,
            3.45,
            0.34,
            [para((body, 13, "muted", False), line_spacing=1.2)],
        )
    for x in (4.46, 8.62):
        add_arrow(slide, x, 2.63, "right", 0.16, "gold")
    add_arrow(slide, 10.82, 3.26, "down", 0.16, "gold")
    for x in (8.46, 4.30):
        add_arrow(slide, x, 3.89, "left", 0.16, "gold")

    entries = [
        (M, "卡片", "申请、调额、冻结与挂失由界面发起；写操作需要密码再认证。"),
        (6.86, "理财", "申购与赎回由界面发起并逐笔确认；Agent 只查询与解释。"),
    ]
    for x, title, body in entries:
        add_rect(slide, x, 4.78, 5.91, 1.28, "blue_pale", None, radius=0.04)
        add_paras(
            slide,
            x + 0.26,
            4.98,
            5.40,
            0.28,
            [para((title, 17, "ink", True), line_spacing=1.0)],
        )
        add_paras(
            slide,
            x + 0.26,
            5.34,
            5.40,
            0.52,
            [para((body, 13.5, "ink", False), line_spacing=1.3)],
        )
    add_paras(
        slide,
        M,
        6.28,
        CONTENT_W,
        0.34,
        [
            para(
                (
                    "Agent 只读与解释场景计划，不能创建、修改或下单；计划不会自动触发转账。",
                    13.5,
                    "muted",
                    False,
                ),
                line_spacing=1.2,
            )
        ],
    )
    footer(slide, "07 / 10")


def page_08(prs):
    slide = new_slide(prs)
    header(
        slide,
        "08 / 安全设计",
        "权限落实到账户 对象 金额与期限",
        "登录身份与资金权限分离；密码、Cookie 与确认凭证都不进入模型。",
    )
    add_framed_picture(slide, KIT / "details" / "transfer-confirmation.png", M, 2.08, 3.46, 4.30)
    caption(slide, M, 6.46, 3.46, "后端确认摘要：对象 金额 费用与余额预览", PP_ALIGN.CENTER, 12)

    layers = [
        ("身份层", "登录会话与资金权限分离；密码再认证只发生在界面，模型看不到凭据。"),
        ("确认层（默认）", "逐笔人工确认，挑战绑定当前会话、原操作与准确参数；关闭窗口不会执行。"),
        ("预授权层", "限定账户、收款对象、单笔与累计额度、最低保留余额与期限；执行时再次复核。"),
    ]
    y = 2.08
    for index, (title, body) in enumerate(layers):
        add_rect(slide, 4.40, y, 8.37, 1.02, "white", "line", 0.75, radius=0.04)
        add_rect(slide, 4.40, y, 0.05, 1.02, ["teal", "gold", "ink"][index])
        add_paras(
            slide,
            4.68,
            y + 0.18,
            2.3,
            0.30,
            [para((title, 17, "ink", True), line_spacing=1.0)],
        )
        add_paras(
            slide,
            6.60,
            y + 0.20,
            6.00,
            0.62,
            [para((body, 14, "ink", False), line_spacing=1.3)],
        )
        y += 1.14

    form_crop = (330 / 1440, (1440 - 950) / 1440, 380 / 1200, (1200 - 550) / 1200)
    add_framed_picture(slide, KIT / "screenshots" / "14-authorization-scope.png", 4.40, 5.46, 4.62, 1.20, crop=form_crop)
    caption(slide, 4.40, 6.70, 4.62, "预授权表单的额度、最低余额与截止时间（截图原文）", size=11.5)

    add_rect(slide, 9.28, 5.46, 3.49, 1.20, "teal_pale", None, radius=0.04)
    add_paras(
        slide,
        9.50,
        5.64,
        3.05,
        0.88,
        [
            para(("结果未知 → 保留预占", 15, "ink", True), line_spacing=1.1),
            para(
                (
                    "按原操作编号核对，避免重复扣款；模型不能执行批准。",
                    12.5,
                    "ink",
                    False,
                ),
                line_spacing=1.25,
                space_before=6,
            ),
        ],
    )
    footer(slide, "08 / 10")


def page_09(prs):
    slide = new_slide(prs)
    header(
        slide,
        "09 / 验证证据",
        "验证记录覆盖回归 环境与人工流程",
        "区分自动化回归、依赖环境复现与人工流程审核。",
    )
    metrics = [
        (M, "235", "项", "自动化回归通过", ("Windows · Python 3.14.6", "含 1 项前端浏览器冒烟"), "teal_pale"),
        (4.72, "39", "项", "新建环境定向测试", ("Python 3.13 按锁文件安装", "pip check 通过"), "gold_pale"),
        (8.88, "10", "项", "人工流程审核通过", ("MAN-01 至 MAN-10", "项目负责人确认"), "blue_pale"),
    ]
    for x, number, unit, title, body, fill in metrics:
        add_rect(slide, x, 2.06, 3.89, 1.78, fill, None, radius=0.04)
        add_paras(
            slide,
            x + 0.24,
            2.26,
            3.40,
            0.62,
            [
                para(
                    [
                        (number, 42, "ink", True),
                        ("  " + unit, 16, "ink", True),
                    ],
                    line_spacing=1.0,
                )
            ],
        )
        add_paras(
            slide,
            x + 0.24,
            3.00,
            3.40,
            0.30,
            [para((title, 15, "ink", True), line_spacing=1.0)],
        )
        add_paras(
            slide,
            x + 0.24,
            3.34,
            3.40,
            0.46,
            [
                para((line, 12, "muted", False), line_spacing=1.2)
                for line in (body if isinstance(body, tuple) else (body,))
            ],
        )

    covers = [
        (M, 4.14, "资金一致性", "双边账本、原子转账与余额校验"),
        (3.47, 4.14, "权限隔离", "跨用户隔离、本人对象归属校验"),
        (M, 5.30, "失败恢复", "保留预占，按原编号核对"),
        (3.47, 5.30, "界面流程", "逐笔确认、预授权范围与再认证"),
    ]
    for x, y, title, body in covers:
        add_rect(slide, x, y, 2.75, 1.02, "white", "line", 0.75, radius=0.05)
        add_paras(
            slide,
            x + 0.20,
            y + 0.18,
            2.35,
            0.28,
            [para((title, 14.5, "ink", True), line_spacing=1.0)],
        )
        add_paras(
            slide,
            x + 0.20,
            y + 0.52,
            2.35,
            0.36,
            [para((body, 11.5, "muted", False), line_spacing=1.2)],
        )

    add_framed_picture(slide, KIT / "diagrams" / "context-routing.png", 6.55, 4.14, 6.22, 1.58)
    caption(
        slide,
        6.55,
        5.82,
        6.22,
        "补充机制：每轮按预算装配上下文并回填后端状态",
        PP_ALIGN.CENTER,
    )
    add_rect(slide, M, 6.48, CONTENT_W, 0.42, "ink", None, radius=0.08)
    add_paras(
        slide,
        M + 0.24,
        6.58,
        CONTENT_W - 0.48,
        0.26,
        [
            para(
                (
                    "历史真实模型样本另有记录（35/36 结构检查通过，1 条因输出预算被保护拦截且无扣款），"
                    "不换算为模型业务成功率。",
                    12.5,
                    "white",
                    True,
                ),
                line_spacing=1.0,
            )
        ],
    )
    footer(slide, "09 / 10")


def page_10(prs):
    slide = new_slide(prs)
    header(
        slide,
        "10 / 收束",
        "以可解释查询和受控执行连接银行代办",
    )
    lines = [
        "查询有依据 · 对象可澄清",
        "执行有边界 · 状态可恢复",
        "本地模拟闭环 · 可运行可复核",
    ]
    y = 2.42
    for item in lines:
        add_rect(slide, M + 0.02, y + 0.14, 0.20, 0.04, "gold")
        add_paras(
            slide,
            M + 0.40,
            y,
            6.20,
            0.44,
            [para((item, 24, "ink", True), line_spacing=1.1)],
        )
        y += 0.94
    add_paras(
        slide,
        M + 0.40,
        5.36,
        6.10,
        0.70,
        [
            para(
                (
                    "AI-Banking 把金融意图落实到可解释查询、明确对象、受控执行与可核对回执。",
                    15,
                    "muted",
                    False,
                ),
                line_spacing=1.35,
            ),
            para(
                (
                    "配套源码、部署说明与验证记录，便于评审运行与复核。",
                    15,
                    "muted",
                    False,
                ),
                line_spacing=1.35,
                space_before=6,
            ),
        ],
    )

    crop = (300 / 1440, (1440 - 970) / 1440, 270 / 1000, (1000 - 520) / 1000)
    add_picture(
        slide,
        KIT / "screenshots" / "02-overview.png",
        6.95,
        2.70,
        5.80,
        2.16,
        crop=crop,
        alpha=26,
    )
    add_rect(slide, 6.95, 2.70, 5.80, 2.16, None, "line", 0.75)
    caption(
        slide,
        6.95,
        5.00,
        5.80,
        "完整工作台与配套技术 / 安全 / 部署材料",
        PP_ALIGN.CENTER,
    )
    footer(slide, "答辩结束", note="AI Banking")


def appendix_a(prs):
    slide = new_slide(prs)
    header(
        slide,
        "附录 A · 备选附页",
        "每轮上下文装配与预算保护",
        "上下文按预算装配；预算不足时停止新增操作，不继续冒险执行。",
    )
    add_framed_picture(slide, KIT / "diagrams" / "context-routing.png", 2.17, 2.06, 9.00, 2.26)
    caption(slide, 2.17, 4.38, 9.00, "每轮装配：系统约束与工具 → 必要时摘录旧轮 → 调用前再查预算", PP_ALIGN.CENTER)

    notes = [
        ("最近 12 轮原文", "更早回合压成带轮次编号、最初目标与省略数量的确定性摘录。"),
        ("工具结果截断", "超过 MODEL_TOOL_OUTPUT_CHARS 时截断，并提示用户分页查询。"),
        ("预算不足即停止", "返回 CONTEXT_BUDGET_EXCEEDED，本轮不执行新操作并保留检查点。"),
    ]
    for (title, body), x in zip(notes, (M, 4.72, 8.88)):
        add_rect(slide, x, 4.86, 3.89, 1.16, "white", "line", 0.75, radius=0.05)
        add_paras(
            slide,
            x + 0.22,
            5.04,
            3.45,
            0.28,
            [para((title, 16, "ink", True), line_spacing=1.0)],
        )
        add_paras(
            slide,
            x + 0.22,
            5.40,
            3.45,
            0.48,
            [para((body, 12.5, "muted", False), line_spacing=1.25)],
        )
    add_paras(
        slide,
        M,
        6.24,
        CONTENT_W,
        0.34,
        [
            para(
                (
                    "token 为本地启发式估算，摘要为确定性摘录而非生成式；未做供应商 tokenizer 校准。",
                    12.5,
                    "muted",
                    False,
                ),
                line_spacing=1.2,
            )
        ],
    )
    footer(slide, "附录 A")


def appendix_b(prs):
    slide = new_slide(prs)
    header(
        slide,
        "附录 B · 备选附页",
        "390×844 浏览器视口示例",
        "用于说明窄屏布局，不代表真实手机设备验收。",
    )
    add_framed_picture(slide, KIT / "screenshots" / "13-workflow-mobile.png", M, 2.06, 2.10, 4.32)
    caption(slide, M, 6.44, 2.10, "390×844 视口", PP_ALIGN.CENTER, 12)

    add_paras(
        slide,
        3.20,
        2.20,
        9.57,
        0.70,
        [para(("该图是浏览器视口截图，不是真实手机设备测试结果", 22, "ink", True), line_spacing=1.15)],
    )
    bullets = [
        "界面为同源原生 Web，没有独立移动端 App；同一套页面适配窄屏。",
        "820 与 390 宽度已确认无横向溢出；未做真实手机设备验收，仍属待办事项。",
        "工作流、确认与预授权流程在窄屏下同样需要逐笔确认，不做简化。",
    ]
    y = 3.24
    for item in bullets:
        add_rect(slide, 3.24, y + 0.12, 0.16, 0.03, "gold")
        add_paras(
            slide,
            3.54,
            y,
            9.20,
            0.60,
            [para((item, 16, "ink", False), line_spacing=1.3)],
        )
        y += 0.86
    add_rect(slide, 3.20, 5.86, 9.57, 0.62, "blue_pale", None, radius=0.05)
    add_paras(
        slide,
        3.44,
        6.00,
        9.10,
        0.34,
        [
            para(
                (
                    "其余运行原图与前端浏览器冒烟用例见交付材料，验收口径以项目状态记录为准。",
                    13.5,
                    "ink",
                    False,
                ),
                line_spacing=1.2,
            )
        ],
    )
    footer(slide, "附录 B")


def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def apply_speaker_notes(prs):
    """把 ppt-kit/slides.json 的逐页讲稿写入备注，便于答辩时对照。"""
    source = ROOT / "ppt-kit" / "slides.json"
    notes = []
    if source.exists():
        import json

        notes = [item.get("speaker", "") for item in json.loads(source.read_text(encoding="utf-8"))]
    appendix = [
        "附录A：每轮先组装系统约束、工具与状态，历史超过12轮或预算紧张时摘录旧轮；"
        "工具结果过长截断并提示分页。预算不足时停止新增操作，不继续冒险执行。",
        "附录B：这是390×844的浏览器视口截图，不是真实手机设备测试结果；"
        "界面为同源原生Web，窄屏下流程与确认同样逐笔确认。",
    ]
    for index, slide in enumerate(prs.slides):
        if index < len(notes):
            set_notes(slide, notes[index])
        elif index - len(notes) < len(appendix):
            set_notes(slide, appendix[index - len(notes)])


def build():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    prs.core_properties.title = "AI Banking 答辩演示"
    prs.core_properties.subject = "竞赛答辩 · 10 页正文与 2 页备选附录"
    prs.core_properties.keywords = "AI-Banking, 模拟银行, 受控操作, 答辩"

    page_01(prs)
    page_02(prs)
    page_03(prs)
    page_04(prs)
    page_05(prs)
    page_06(prs)
    page_07(prs)
    page_08(prs)
    page_09(prs)
    page_10(prs)
    appendix_a(prs)
    appendix_b(prs)

    apply_speaker_notes(prs)

    prs.save(OUT)
    print(f"Created {OUT} with {len(prs.slides)} slides")


if __name__ == "__main__":
    build()
